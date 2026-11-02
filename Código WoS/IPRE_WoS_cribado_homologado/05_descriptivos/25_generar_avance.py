"""Genera un informe DOCX de avance WoS desde el JSON y los gráficos.

Uso: py -3 05_descriptivos/25_generar_avance.py --resultados resultados_cribado_wos_nueva
El documento solo describe resultados automáticos y tareas pendientes.
"""
from __future__ import annotations
import argparse
import csv
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
HERE=Path(__file__).resolve().parent

def fmt(x): return f'{int(x):,}'.replace(',','.')
def pp(n,d):return f'{100*n/d:.1f}'.replace('.',',')+' %' if d else '—'

def main():
    p=argparse.ArgumentParser(description='Genera informe Word con gráficos de avance WoS.')
    p.add_argument('--resultados',default='resultados_cribado_wos')
    p.add_argument('--salida',default=None,help='Carpeta de descriptivos si se generaron fuera de 05_descriptivos.')
    a=p.parse_args()
    dest=Path(a.salida) if a.salida else HERE
    if not dest.is_absolute():dest=ROOT/dest
    data_file=dest/'Resultados_descriptivos_amplio_y_v9_WoS.json'
    if not data_file.exists():raise FileNotFoundError('Primero ejecute 24_generar_graficos.py: '+str(data_file))
    d=json.loads(data_file.read_text(encoding='utf-8'))
    requested=Path(a.resultados).name
    if d['carpeta_resultados']!=requested:
        raise ValueError(f'Los gráficos fueron generados para {d["carpeta_resultados"]}, pero solicitó {requested}. Vuelva a ejecutar 24_generar_graficos.py.')
    figures=dest/'graficos'
    required=['01_flujo_del_estudio','02_versiones_del_cribado','03_cierre_de_TC',
       '04_publicaciones_por_ano','05_menciones_de_procesos','06_menciones_de_equipos',
       '07_focos_en_cada_base','08_muestra_pendiente_revision_humana']
    if any(not (figures/(x+'.png')).exists() for x in required):
        raise FileNotFoundError('Faltan gráficos. Ejecute primero 24_generar_graficos.py.')
    out=dest/'informe';out.mkdir(parents=True,exist_ok=True)
    file=out/'Informe_avance_IPRE_WoS_preliminar.docx'
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.72);section.bottom_margin=Inches(.67)
    section.left_margin=Inches(.78);section.right_margin=Inches(.78)
    section.header_distance=Inches(.32);section.footer_distance=Inches(.32)
    for style in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        s=doc.styles[style];s.font.name='Arial';s.font.color.rgb=RGBColor(22,38,49)
        s.paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].font.size=Pt(10.0)
    doc.styles['Normal'].paragraph_format.line_spacing=1.09
    doc.styles['Title'].font.size=Pt(21)
    doc.styles['Subtitle'].font.size=Pt(11)
    doc.styles['Heading 1'].font.size=Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(14)
    doc.styles['Heading 2'].font.size=Pt(11)
    doc.styles['Caption'].font.size=Pt(8.8)
    doc.styles['Caption'].font.italic=False
    head=section.header.paragraphs[0];head.text='IPRE   |   Web of Science   |   Minería a cielo abierto'
    head.runs[0].font.size=Pt(8)
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('Avance preliminar   ·   ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    props=doc.core_properties;props.title='IPRE — Avance de cribado WoS';props.author='Proyecto IPRE';
    props.subject='Resultados computacionales de Web of Science y revisión humana pendiente'

    def para(text='',boldlead=None):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if boldlead and text.startswith(boldlead):p.add_run(boldlead).bold=True;p.add_run(text[len(boldlead):])
        else:p.add_run(text)
        return p
    def heading(text):doc.add_heading(text,level=1)
    def subheading(text):doc.add_heading(text,level=2)
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,s in enumerate(headers):t.rows[0].cells[i].text=str(s)
        for values in rows:
            cells=t.add_row().cells
            for cell,val in zip(cells,values):cell.text=str(val)
        if widths:
            for row in t.rows:
                for i,cell in enumerate(row.cells):cell.width=Inches(widths[i])
        for ri,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for c in row.cells:
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=c._tc.get_or_add_tcPr()
                if ri==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDE8EE');cp.append(sh)
                for pr in c.paragraphs:
                    pr.paragraph_format.space_after=Pt(0);pr.paragraph_format.space_before=Pt(0)
                    for r in pr.runs:r.font.size=Pt(8.5);r.bold=(ri==0)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
    def fig(name,caption,width=6.75):
        par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before=Pt(5);par.paragraph_format.space_after=Pt(2)
        par.paragraph_format.keep_with_next=True
        par.add_run().add_picture(str(figures/(name+'.png')),width=Inches(width))
        cap=doc.add_paragraph(style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.keep_together=True;cap.add_run(caption)
    def toptags(which,field,n=6):
        dd=d[which][field];return [(k,v) for k,v in list(dd.items())[:n]]
    def names(table):
        return [(k.replace('_',' ').capitalize(),fmt(v)) for k,v in table]
    n=d['base_amplia']['N'];core=d['in_propuestos_computacionalmente']['N'];f1=d['fase01']
    ybase=d['base_amplia'];ycore=d['in_propuestos_computacionalmente']
    end='1 de noviembre de 2026'
    doc.add_paragraph('Avance del cribado de publicaciones de Web of Science',style='Title')
    doc.add_paragraph('Replicación del flujo IPRE_codigo · Resultados descriptivos y estado de validación',style='Subtitle')
    para(f'Fuente: Web of Science Core Collection  |  Corte del estudio: {end}')
    para('Estado del informe: resultados generados automáticamente con los motores del proyecto IPRE. No se ha incorporado aún la revisión humana independiente de WoS ni una resolución semántica validada de las publicaciones.')

    heading('1. Propósito y alcance')
    para('Este avance describe la aplicación del flujo de cribado desarrollado en IPRE_codigo a registros bibliográficos de Web of Science (WoS). El interés de la investigación es caracterizar literatura de minería a cielo abierto y distinguir estudios que evalúan procesos, tareas o equipos operacionales de aquellos orientados a otros problemas, manteniendo atención sobre las actividades y los equipos auxiliares.')
    para('La extracción de WoS se realizó mediante diez exportaciones web manuales incrementales. Ese cambio en el mecanismo de recuperación no modifica por sí solo las reglas posteriores de selección. La normalización y el filtro amplio corresponden a la Fase 1, mientras que las etapas v6, v7, v8 y v9 se aplicaron sobre la misma base amplia de WoS.')
    table(['Elemento','Alcance documentado'],[
       ['Base de datos','Web of Science Core Collection'],
       ['Método de extracción','Exportación manual incremental; sin API'],
       ['Unidad de seguimiento','Registro bibliográfico identificado por wos_id'],
       ['Análisis de contenido','Título, resumen y metadatos conservados en las matrices'],
       ['Estado de decisiones','Propuestas automáticas, todavía no dictámenes humanos finales'],
    ],[2.15,4.75])
    fig('01_flujo_del_estudio','Figura 1. Universo WoS, filtro amplio y núcleo operacional propuesto. Los 168 candidatos IN y 2.070 OUT se incluyen dentro de los 2.238 registros de la base amplia.',width=6.9)

    heading('2. Extracción, normalización y primer filtro')
    table(['Etapa','Registros','Criterio de interpretación'],[
      ['Registros únicos',fmt(f1['total_uid']),'Total de UID exportados y normalizados'],
      ['Sin resumen',fmt(f1['sin_resumen']),'Conservados por separado para control documental'],
      ['Excluidos por filtro amplio',fmt(f1['excluidos_filtro_amplio']),'Fuera del conjunto candidato a cribado operativo'],
      ['Base amplia WoS',fmt(n),'Población de referencia para v6–v9'],
    ],[2.25,1.05,3.6])
    para('La base amplia concentra los registros que superaron los controles iniciales de la Fase 1. Los documentos sin resumen y los excluidos por filtro amplio no se han eliminado del proyecto: se conservan en archivos diferenciados para eventuales verificaciones, análisis de sensibilidad y trazabilidad del proceso de selección.')
    para('Se registraron dos alertas documentales y tres grupos de posibles coincidencias bibliográficas con identificadores diferentes. Estos casos requieren comprobación en la fuente antes de cualquier decisión de retractación, deduplicación adicional o integración entre bases.')

    heading('3. Criterios aplicados y versiones del cribado')
    para('El motor de v6 examina el contexto de minería a cielo abierto y señales relacionadas con proceso/equipo, decisión o evaluación y resultado operacional. v7 resuelve automáticamente las fichas dudosas de la etapa previa. v8 aplica el criterio calibrado originalmente en IPRE_codigo sin importar las respuestas humanas del corpus OpenAlex. v9 aplica reglas de cierre a los casos TC de v8, sin realizar una nueva lectura semántica de texto completo.')
    table(['Etapa','IN','OUT','DUDOSO / TC','Estado'],[
      ['v6',fmt(d['v6'].get('IN_FUERTE',0)+d['v6'].get('IN_SECUNDARIO',0)),fmt(d['v6'].get('OUT',0)),fmt(d['v6'].get('DUDOSO',0)),'Primer cribado por dimensiones'],
      ['v7',fmt(d['v7'].get('IN',0)),fmt(d['v7'].get('OUT',0)),'0','Resolución automática de dudosos'],
      ['v8',fmt(d['v8'].get('IN',0)),fmt(d['v8'].get('OUT',0)),fmt(d['v8'].get('TC',0)),'Propuesta con criterio ajustado'],
      ['v9',fmt(d['v9_solo_reglas'].get('IN',0)),fmt(d['v9_solo_reglas'].get('OUT',0)),'0','Solo reglas, sin validación humana'],
    ],[.62,.63,.63,1.1,3.92])
    fig('02_versiones_del_cribado','Figura 2. Distribución de propuestas sobre los mismos 2.238 registros en cada versión. Los DUDOSO de v6 se visualizan junto con TC, pero no equivalen exactamente a la categoría TC de v8.',width=6.9)
    para('Las diferencias entre versiones reflejan cambios en las reglas y en la forma de evaluar la evidencia bibliográfica. No deben interpretarse como nuevas publicaciones agregadas, ni como desacuerdos observados frente a una revisión humana independiente.')

    heading('4. Tratamiento de los 917 casos TC en v8')
    c=d['tc_v8_resueltos_solo_reglas']
    table(['Salida computacional de v9','Registros'],[
       ['IN con evidencia en título o resumen',fmt(c.get('inclusion_con_evidencia',0))],
       ['OUT por foco temático diferente',fmt(c.get('exclusion_por_foco',0))],
       ['OUT por evidencia documental insuficiente',fmt(c.get('exclusion_documental',0))],
       ['Total de casos TC examinados por reglas',fmt(sum(c.values()))],
    ],[5.18,1.72])
    fig('03_cierre_de_TC','Figura 3. Disposición automática de casos TC de v8. «OUT documental» indica evidencia insuficiente en la ficha disponible y no implica irrelevancia comprobada del texto completo.',width=6.75)
    para(f'El cierre v9 propone {fmt(core)} IN y {fmt(n-core)} OUT dentro de la base amplia. Estas cifras combinan las decisiones conservadas de v8 y la aplicación de reglas sobre los {fmt(sum(c.values()))} TC. El número final de publicaciones elegibles debe determinarse después de las revisiones humanas y los controles bibliográficos pendientes.')

    heading('5. Distribución temporal de las publicaciones')
    fig('04_publicaciones_por_ano','Figura 4. Registros por año editorial entre 2000 y 2025. El año 2026 se separa por estar incompleto en el corte de extracción.',width=6.8)
    table(['Periodo','Base amplia','IN propuestos'],[
       ['Antes de 2000',fmt(ybase['antes_2000']),fmt(ycore['antes_2000'])],
       ['2000–2025',fmt(ybase['2000_2025']),fmt(ycore['2000_2025'])],
       ['2026 (parcial)',fmt(ybase['2026_parcial']),fmt(ycore['2026_parcial'])],
       ['Año no disponible',fmt(ybase['sin_ano']),fmt(ycore['sin_ano'])],
    ],[3.55,1.62,1.73])
    para('La variable year corresponde al año editorial conservado en la exportación. El registro también incluye campos de fecha de publicación y Early Access, que pueden no coincidir. Antes de realizar comparaciones bibliométricas temporales definitivas se deben verificar estas diferencias y la cobertura de las consultas.')

    heading('6. Procesos identificados por menciones automáticas')
    fig('05_menciones_de_procesos','Figura 5. Etiquetas automáticas de procesos mineros. Un mismo trabajo puede activar varias etiquetas y ninguna prueba por sí sola cuál es el proceso central estudiado.',width=6.8)
    for k,v in toptags('base_amplia','auto_process_tags',5):
        para(f'{k.replace("_"," ").capitalize()}: {fmt(v)} registros con coincidencia en la base amplia.')
    para('El análisis de los procesos auxiliares requiere complementar estas menciones con lectura de título y resumen, distinguiendo la tarea estudiada, el tipo de intervención o decisión y los resultados que el artículo efectivamente evalúa. Un trabajo de geotecnia o tronadura puede mencionar maquinaria sin medir su desempeño.')

    heading('7. Equipos presentes en la literatura')
    fig('06_menciones_de_equipos','Figura 6. Presencia automática de familias de equipos en la base amplia y entre las inclusiones propuestas. Las familias no son excluyentes y pueden incluir menciones incidentales.',width=6.8)
    pairs=d['base_amplia']['equipment_title_tags'];coretitle=d['in_propuestos_computacionalmente']['equipment_title_tags']
    keys=list(pairs)[:5]
    table(['Equipo mencionado en título','Base amplia','IN propuesto'],[
        [k.replace('_',' ').capitalize(),fmt(pairs[k]),fmt(coretitle.get(k,0))] for k in keys
    ],[3.75,1.46,1.69])
    para('Las coincidencias en el título sirven como señal adicional, pero no sustituyen la evaluación del contenido del estudio. Los conteos bajos de ciertas categorías auxiliares tampoco demuestran por sí solos que esas actividades no estén investigadas: podrían aparecer descritas por funciones, sinónimos o nombres alternativos.')

    heading('8. Enfoques temáticos detectados')
    fig('07_focos_en_cada_base','Figura 7. Porcentaje de publicaciones con cada etiqueta temática automática. Los porcentajes son independientes y no deben sumarse como una partición de temas.',width=6.8)
    para('Este gráfico permite observar el vocabulario disponible para analizar planificación, desempeño operacional, seguridad, mantenimiento y automatización. Una etiqueta de desempeño identifica señales textuales, pero no acredita que el trabajo haya medido un resultado atribuible a una tarea o equipo. La clasificación sustantiva debe desarrollarse con revisión documental.')

    heading('9. Revisión humana de v6: estado y procedimiento')
    fig('08_muestra_pendiente_revision_humana','Figura 8. Diseño de la muestra ciega de v6: 50 candidatos IN, 50 OUT y 30 DUDOSO. No corresponde a resultados de calibración ni a decisiones humanas completadas.',width=6.8)
    para('La muestra consta de 130 publicaciones. La revisión independiente debe realizarse sin consultar la clasificación automática de cada ficha, utilizando identificadores neutrales y registrando la decisión IN, OUT o TC junto con su justificación textual y, cuando corresponda, la necesidad de revisar el documento completo.')
    para('Al finalizar, se contrastarán los juicios humanos con las propuestas automáticas, se documentarán los desacuerdos y se decidirá si corresponde modificar los criterios antes de consolidar las fases posteriores. Hasta completar este trabajo, no es apropiado reportar precisión, acuerdo interevaluador ni tasa de error de las decisiones automáticas.')

    heading('10. Reproducción y control de integridad')
    para('Los programas de la carpeta 05_descriptivos únicamente leen las matrices y sus manifiestos. No alteran los UID WoS, las decisiones de v6 a v9, los resultados de la Fase 1 ni las plantillas de revisión humana. El resumen descriptivo conserva el SHA-256 de la matriz v9 utilizada, y el generador comprueba la coincidencia con el manifiesto del cribado.')
    table(['Componente','Archivo de trabajo'],[
      ['Cribado principal','00_ejecutar_cribado_wos.py'],
      ['Generador de gráficos','05_descriptivos/24_generar_graficos.py'],
      ['Generador de informe','05_descriptivos/25_generar_avance.py'],
      ['Ejecución descriptiva conjunta','05_descriptivos/00_generar_descriptivos.py'],
      ['Estadísticas y trazabilidad','05_descriptivos/Resultados_descriptivos_amplio_y_v9_WoS.json'],
      ['Gráficos','05_descriptivos/graficos/'],
      ['Informe Word','05_descriptivos/informe/Informe_avance_IPRE_WoS_preliminar.docx'],
    ],[2.55,4.35])
    para('El parámetro --resultados permite generar nuevamente los gráficos y este informe sobre cualquier carpeta de salida producida por el cribado principal. Se requiere que esa carpeta incluya RESUMEN_FLUJO_WOS.json y las matrices originales de v6 a v9.')

    heading('11. Limitaciones del avance')
    for msg in [
        'Las 168 inclusiones y 2.070 exclusiones de v9 no han sido validadas mediante un ejercicio humano independiente de WoS.',
        'El criterio v8 fue trasladado desde un piloto previo sobre OpenAlex. No se puede asumir que su desempeño sea idéntico en WoS sin calibración adicional.',
        'Los 917 casos TC de v8 fueron tratados con reglas automatizadas. No se realizaron lecturas semánticas independientes ni lectura de textos completos en esta ejecución.',
        'Persisten alertas de posibles retractaciones y coincidencias bibliográficas entre UID diferentes. Las fechas editoriales y de publicación electrónica requieren revisión antes de análisis temporales definitivos.',
        'Los conteos automáticos de procesos y equipos registran menciones y no el foco central verificado de cada artículo.',
        'No se ha realizado una deduplicación entre WoS, Scopus y OpenAlex. Por lo tanto, los registros de bases distintas no deben sumarse como publicaciones únicas.',
    ]:para('• '+msg)
    heading('12. Próximos pasos')
    for msg in [
        'Completar la evaluación ciega de las 130 fichas de v6 con criterios de inclusión y exclusión documentados.',
        'Contrastar las decisiones humanas y automáticas, registrar los desacuerdos y revisar los criterios aplicados a los casos limítrofes.',
        'Reexaminar los casos TC de v8 y las inclusiones propuestas por v9, priorizando el texto completo cuando el resumen sea insuficiente.',
        'Comprobar las alertas documentales, los grupos con posible duplicación y las diferencias de fecha editorial y Early Access.',
        'Conservar la versión automática como línea base y publicar nuevas matrices solo después de integrar decisiones documentadas.',
        'Homologar identificadores y deduplicar entre fuentes antes de presentar conclusiones bibliométricas consolidadas.',
    ]:para('• '+msg)
    para('Conclusión preliminar: WoS cuenta con una base amplia de 2.238 registros y un subconjunto computacional de 168 propuestas IN para estudiar desempeño operacional. El flujo y las figuras son reproducibles desde archivos locales, mientras que la elegibilidad final de los estudios y las conclusiones sobre brechas en procesos auxiliares permanecen sujetas a revisión.')

    doc.save(file)
    print(file)

if __name__=='__main__':main()
