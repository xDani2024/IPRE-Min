from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import argparse
ROOT = next(p for p in Path(__file__).resolve().parents if (p/'PROCEDENCIA_DEL_PAQUETE.json').is_file())

parser=argparse.ArgumentParser()
parser.add_argument('--entrada',type=Path,default=ROOT/'04_cierre_v9/salida')
parser.add_argument('--salida',type=Path,default=ROOT/'05_descriptivos/salida')
args=parser.parse_args()
args.salida.mkdir(parents=True,exist_ok=True)
exec(compile((Path(__file__).parent/'_descriptivos.py').read_text(encoding='utf-8'),str(Path(__file__).parent/'_descriptivos.py'),'exec'))
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
DOCX=SALIDA/'Avance_IPRE_Scopus_v9_PRELIMINAR.docx'
doc=Document();sec=doc.sections[0]
sec.top_margin=Inches(.66);sec.bottom_margin=Inches(.66)
sec.left_margin=Inches(.83);sec.right_margin=Inches(.83)
for key in ['Normal','Title','Heading 1','Heading 2','Caption']:
    style=doc.styles[key];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(9.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Title'].font.size=Pt(18)
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Caption'].font.size=Pt(8)
doc.core_properties.title='Avance descriptivo preliminar IPRE — Scopus v9'
doc.core_properties.subject='Análisis descriptivo de publicaciones Scopus, resultados computacionales pendientes de validación'
header=sec.header.paragraphs[0];header.text='IPRE   |   Minería a cielo abierto   |   Scopus'
header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Avance preliminar  |  19 de octubre de 2026  |  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def par(s):return doc.add_paragraph(s)
def hd(s):doc.add_heading(s,level=1)
def tabla(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.autofit=True
    for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
    for r in rows:
        cs=t.add_row().cells
        for c,v in zip(cs,r):c.text=str(v)
    doc.add_paragraph()
def graf(n,leyenda,width=6.65):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIG/(n+'.png')),width=Inches(width))
    c=doc.add_paragraph(leyenda,style='Caption');c.alignment=WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Análisis descriptivo preliminar de Scopus',style='Title')
par('IPRE — Minería a cielo abierto | Fase 05 | 19 de octubre de 2026')
par('Esta etapa aplica a Scopus el esquema descriptivo de Joaquín. Los resultados derivan de la matriz computacional v9 y no deben interpretarse como una revisión humana validada.')
hd('1. Universo de análisis y alcance')
tabla(['Resultado','Publicaciones'],[
    ['Coincidencias de las diez búsquedas',fmt(etapa01["coincidencias"])],['Publicaciones únicas tras deduplicación',fmt(etapa01["unicos"])],
    ['Exclusiones documentales',fmt(etapa01["documentales"])],['Excluidas por filtro amplio',fmt(etapa01["excluidos_amplios"])],
    ['Base amplia Scopus',fmt(len(registros))],['Núcleo operacional IN propuesto',fmt(len(nucleo))],
    ['Fuera del núcleo, permanecen en base amplia',fmt(conteos_v9["OUT"])]])
par('Las cifras de extracción provienen del manifiesto de Fase 1 incluido. La matriz de Fase 4 permite recalcular directamente las cifras de base amplia y núcleo propuesto. Las exclusiones documentales y temáticas no se vuelven a ejecutar en este notebook.')
graf('01_flujo_del_estudio','Figura 1. Flujo del corpus Scopus. Todas las inclusiones del núcleo son provisionales.',6.2)

doc.add_page_break()
hd('2. Evolución de la clasificación')
par(f'Las versiones v6–v9 se presentan sobre las mismas {fmt(len(registros))} publicaciones. La v9 es un pre-cierre computacional sin transferir auditorías o juicios humanos históricos de OpenAlex.')
tabla(['Versión','IN','OUT','TC / DUDOSO'],[
    [k, data[k].get('IN',0),data[k].get('OUT',0),data[k].get('TC',0)] for k in ['v6','v7','v8','v9']])
graf('02_versiones_del_cribado','Figura 2. Evolución de propuestas, con idéntico denominador.',6.25)
par(f'La v9 evaluó {len(tc_previos):,} TC previos: {disposiciones_tc.get("inclusion_con_evidencia",0)} IN con evidencia computacional, {disposiciones_tc.get("exclusion_por_foco",0)} OUT por foco y {disposiciones_tc.get("exclusion_documental",0)} OUT documentales.'.replace(',','.'))
graf('03_cierre_de_TC','Figura 3. Resolución automática de TC. OUT documental no demuestra irrelevancia del texto completo.',6.25)

hd('3. Distribución temporal y etiquetas descriptivas')
par('La serie gráfica usa 2000–2025, igual que el análisis original. Los años 2026 y posteriores no se eliminan: se informan separadamente debido al corte temporal no restringido de la consulta de origen.')
tabla(['Grupo','Antes de 2000','2000–2025','Año 2026','Posteriores a 2026','Sin año'],[
    [lab,*(data[k][i] for i in ['antes_2000','2000_2025','2026_parcial','posteriores_2026','sin_ano'])] for k,lab in [('base_amplia','Base amplia'),('nucleo_propuesto_v9','Núcleo IN')]])
graf('04_publicaciones_por_ano','Figura 4. Evolución 2000–2025. El año 2026 es parcial.',6.3)
par('Los diccionarios de Joaquín detectan menciones de procesos, equipos y focos. Las categorías pueden superponerse dentro de un artículo y no confirman el objeto principal. La comparación visual está sujeta a los límites del vocabulario utilizado.')
graf('05_menciones_de_procesos','Figura 5. Menciones automáticas de procesos, no temas centrales verificados.',6.5)
graf('06_menciones_de_equipos','Figura 6. Menciones automáticas de equipos, incluyendo posibles auxiliares.',5.2)
graf('07_focos_en_cada_base','Figura 7. Porcentajes por foco dentro de cada grupo; distintos denominadores.',5.2)

hd('4. Validación pendiente y próximos pasos')
par('Joaquín dispone de 40 fichas históricas con juicio humano en OpenAlex. Estas no se trasladaron a Scopus. A la fecha, los resultados de Scopus no incluyen auditoría humana independiente ni lectura de texto completo en el cierre v9.')
graf('08_estado_validacion_pendiente','Figura 8. Trabajos de revisión propuestos, no decisiones humanas realizadas.',4.95)
par(f'Prioridades: (1) revisar las {len(nucleo)} propuestas IN, incluidas las {man["TC_resultado_computacional"].get("IN",0)} nuevas desde TC; (2) evaluar 120 OUT seleccionados y controlar los excluidos amplios; (3) resolver alertas de deduplicación, retractaciones y paratextos de Fase 1; (4) definir objeto central y tipo de evidencia por publicación para contrastar las hipótesis sobre equipos auxiliares.')
par('La similitud de recuentos Scopus y OpenAlex no implica igualdad de publicaciones ni equivalencia de validación. Los resultados son preliminares y deben mantenerse como una cohorte independiente.')

# Quitar bordes decorativos heredados del estilo de título.
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for paragraph in doc.paragraphs:
    for border in list(paragraph._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
# Mantener las tablas pequeñas unidas y repetir sus encabezados.
for table in doc.tables:
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next=i<len(table.rows)-1
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        element=OxmlElement('w:'+edge);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');borders.append(element)
    table._tbl.tblPr.append(borders)
doc.save(DOCX)
print('Informe Word:',DOCX)
manifest={'fuente':'Scopus','base_amplia':len(registros),'nucleo_in_propuesto':len(nucleo),'graficos':len(list(FIG.glob('*.png'))),'validacion_humana_completada':False,'artefactos_sha256':{str(p.relative_to(SALIDA)):sha256(p) for p in SALIDA.rglob('*') if p.is_file() and p.suffix in ('.png','.json','.docx') and p.name!='MANIFIESTO_FASE05.json'}}
(SALIDA/'MANIFIESTO_FASE05.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
