"""Genera el avance del estudio con resultados y figuras conservados."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = ROOT / 'documentacion'
OUT.mkdir(exist_ok=True)
D = json.loads((HERE / 'Resultados_descriptivos_amplio_y_v9.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.27)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(6)
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.10
doc.styles['Title'].font.size = Pt(21)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(3)
doc.styles['Heading 2'].font.size = Pt(11.5)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
head = sec.header.paragraphs[0]
head.text = 'IPRE   Minería a cielo abierto   Avance del estudio'
head.runs[0].font.size = Pt(8.5)
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Joaquín León  |  8 de octubre de 2026  |  ').font.size = Pt(8.5)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
foot._p.append(field)
doc.core_properties.author = 'Joaquín León'
doc.core_properties.title = 'Avance bibliométrico en minería a cielo abierto'
doc.core_properties.subject = 'Preguntas e hipótesis de investigación y resultados preliminares'

def para(text, bold=False, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(text).bold = bold
    return p

def heading(text, new=True):
    p = doc.add_heading(text, level=1)
    if new:
        p.paragraph_format.page_break_before = True

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = widths or [6.9 / len(headers)] * len(headers)
    for col, w in zip(t.columns, widths):
        col.width = Inches(w)
    for cell, w, text in zip(t.rows[0].cells, widths, headers):
        cell.width = Inches(w)
        cell.text = text
    for row in rows:
        cells = t.add_row().cells
        for c, w, value in zip(cells, widths, row):
            c.width = Inches(w)
            c.text = str(value)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side, val in [('top', '85'), ('bottom', '85'), ('left', '100'), ('right', '100')]:
                e = OxmlElement('w:' + side)
                e.set(qn('w:w'), val)
                e.set(qn('w:type'), 'dxa')
                mar.append(e)
            pr.append(mar)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + side)
                e.set(qn('w:val'), 'single')
                e.set(qn('w:sz'), '4')
                e.set(qn('w:color'), 'D9D9D9')
                borders.append(e)
            pr.append(borders)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'DDE6ED')
                pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(10)
                    r.bold = i == 0
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def figure(name, caption, width=6.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(HERE / 'graficos' / name), width=Inches(width))
    para(caption, style='Caption')

def code(text, boxed=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(8.5 if boxed else 10)
    if boxed:
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(.12)
        pr = p._p.get_or_add_pPr()
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'F4F5F6')
        pr.append(shade)
        borders = OxmlElement('w:pBdr')
        for side in ['top', 'left', 'bottom', 'right']:
            border = OxmlElement('w:' + side)
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:space'), '6')
            border.set(qn('w:color'), 'D9D9D9')
            borders.append(border)
        pr.append(borders)

doc.add_paragraph('Avance bibliométrico en minería a cielo abierto', style='Title')
doc.add_paragraph('Estado del corpus y resultados preliminares', style='Subtitle')
para('Joaquín León | Profesor guía José Joaquín Jara | 8 de octubre de 2026')
heading('1 Contexto y objetivo', new=False)
para('El estudio busca ordenar la literatura sobre la operación de minas a cielo abierto, con especial interés en los equipos y actividades auxiliares. A partir de lo conversado con el profesor, el objetivo es comparar qué se ha investigado sobre planificación, mantenimiento, asignación de recursos y eficiencia durante la ejecución, para después identificar posibles brechas en los procesos de apoyo.')
para('El profesor propuso mirar la mina como un sistema de procesos. Distinguió los procesos principales, como perforación, tronadura, carguío y transporte; las actividades auxiliares que permiten ejecutarlos; el mantenimiento; y la captura y procesamiento de información. El equipo se estudia dentro de la tarea que cumple, por lo que necesitamos conservar ambos ejes: proceso y equipo.')
para('Comenzamos con búsquedas por nombres de equipos y términos de desempeño. Luego construimos una base más amplia de minería a cielo abierto. En la última revisión contamos con 4.735 referencias y un subconjunto de 419 estudios propuestos para analizar desempeño operacional. Las decisiones y sus motivos quedaron registrados por identificador.')
table(['Resultado actual', 'Estado'], [
    ['Base amplia', '4.735 registros conservados'],
    ['Subconjunto de desempeño operacional', '419 propuestos para incluir y 4.316 fuera de ese subconjunto'],
    ['Casos antes pendientes', '2.399 resueltos: 57 incluidos y 2.342 excluidos'],
    ['Revisión humana realizada', '40 fichas de calibración con sugerencias visibles'],
    ['Revisión humana independiente', 'Pendiente']], widths=[3.0, 3.9])
para('La base amplia sirve para comparar las áreas de investigación. El subconjunto permite profundizar en estudios que evalúan el funcionamiento de equipos o tareas. Estar fuera del subconjunto de desempeño no elimina una publicación de la base general.')
para('Usamos IN para una propuesta de inclusión, OUT para quedar fuera del subconjunto y TC para los casos que necesitaban revisión adicional, habitualmente de texto completo. La última propuesta ya no tiene TC pendientes, aunque conserva las respuestas y motivos de las revisiones anteriores.')

heading('2 Preguntas e hipótesis de investigación')
para('La idea es partir del panorama general y avanzar hacia el estudio de los procesos auxiliares. Las siguientes preguntas e hipótesis orientan ese análisis. Hasta ahora tenemos resultados exploratorios y una base de trabajo; las hipótesis todavía deben contrastarse.')
doc.add_heading('Preguntas de investigación', level=2)
para('1. ¿Cómo se distribuye la literatura de minería a cielo abierto entre diseño y planificación, mantenimiento, procesos principales y procesos auxiliares?')
para('2. En los procesos principales, ¿predominan los estudios de planificación y asignación de tareas y equipos sobre los que evalúan la eficiencia durante la ejecución?')
para('3. En los procesos auxiliares, ¿cuánta investigación aborda la asignación óptima de recursos y el desempeño o eficiencia de las actividades de apoyo?')
doc.add_heading('Hipótesis planteadas por el profesor', level=2)
para('H1. La literatura se concentra en el diseño y la planificación de la operación, el mantenimiento y los procesos principales, especialmente perforación, tronadura, carguío y transporte.')
para('H2. En los procesos principales predominan el diseño y la planificación de corto plazo y la asignación eficiente de equipos y tareas. Existe menos investigación sobre el desempeño operacional y la eficiencia durante la ejecución de las actividades productivas.')
para('H3. Existe escasa literatura sobre la asignación óptima y el desempeño o eficiencia de los procesos auxiliares en minería a cielo abierto.')
doc.add_heading('Cómo vamos a contrastarlas', level=2)
para('Primero clasificaremos la base amplia por proceso, equipo y tipo de problema. Con esa clasificación podremos comparar planificación, asignación y ejecución dentro de los procesos principales, y después revisar en detalle qué se ha estudiado sobre los auxiliares. En estos últimos registraremos además los indicadores y métodos utilizados.')
para('Para las dos primeras hipótesis necesitamos la base amplia. Usar solo los 419 seleccionados por desempeño dejaría fuera, por definición, parte de la literatura de planificación o mantenimiento que queremos comparar. Los conteos de palabras de este avance orientan la revisión, pero todavía no confirman ninguna de las tres hipótesis.')

heading('3 De dónde vienen los datos')
para('Los datos provienen de OpenAlex. El análisis parte de los 9.534 registros disponibles, consolidados a partir de las salidas conservadas del filtro general. Esta es la base que usamos para reproducir el procesamiento y calcular los resultados del avance.')
figure('01_flujo_del_estudio.png', 'Figura 1. Flujo del análisis desde los 9.534 registros disponibles hasta la base amplia y el subconjunto de desempeño operacional.', width=6.65)
para('El filtro general conserva 4.735 referencias relacionadas con el ámbito del estudio y deja fuera 4.799. Sobre las 4.735 aplicamos el criterio de desempeño operacional: 419 quedan propuestos para incluir y 4.316 quedan fuera de ese subconjunto.')
para('Las dos selecciones responden a preguntas distintas. La primera delimita la literatura de minería a cielo abierto; la segunda identifica estudios que evalúan el desempeño de equipos o tareas. Los gráficos comparan la base amplia de 4.735 con los 419 candidatos de desempeño. La base amplia y la matriz actual comparten los mismos identificadores, títulos y resúmenes.')

heading('4 Búsqueda y filtro general')
para('Las primeras búsquedas por equipo recuperaban pocas referencias de algunos auxiliares y mezclaban estudios del equipo con menciones incidentales. Buscar también por función permitió encontrar trabajos sobre drenaje, caminos y riego que no siempre nombran la máquina. Por eso ampliamos la recuperación al contexto de minería a cielo abierto y aplicamos la clasificación temática después.')
table(['Parámetro', 'Configuración utilizada'], [
    ['Fuente', 'OpenAlex API, colección de publicaciones works'],
    ['Campo de búsqueda', 'Título y resumen con expresiones exactas'],
    ['Expresiones de contexto', 'Diez expresiones alternativas de minería a cielo abierto, detalladas en la consulta'],
    ['Filtros técnicos', 'Artículo o revisión; resumen disponible; sin retractación y sin paratexto'],
    ['Año e idioma', 'Sin restricción explícita en la consulta archivada'],
    ['DOI', 'Se conserva cuando está disponible; no es obligatorio'],
    ['Descarga', 'Paginación por cursor y registro para reanudar la ejecución']], widths=[2.1, 4.8])
para('Consulta archivada a https://api.openalex.org/works, expresada como el parámetro filter del código:')
code('''terminos = [
    "open pit mining", "open-pit mining", "open pit mine",
    "open-pit mine", "surface mining", "surface mine",
    "opencast mining", "open-cast mining",
    "open cut mining", "open-cut mining"]
filtro = "title_and_abstract.search.exact:" + "|".join(
    f'"{t}"' for t in terminos)
filtro += ",type:article|review,has_abstract:true"
filtro += ",is_retracted:false,is_paratext:false"''', boxed=True)
para('La barra vertical combina expresiones alternativas; las comas añaden condiciones. La búsqueda se aplica al título y al resumen, sin restricción explícita de año o idioma.')
para('Organizamos el trabajo en dos códigos. El primero descarga y normaliza los metadatos, reconstruye los resúmenes y consolida duplicados por identificador, DOI y similitud bibliográfica, con controles de año y autor. El segundo aplica el filtro amplio y conserva el historial de decisiones. El cruce documentó 18 equivalencias entre identificadores de una misma obra.')
para('No exigimos DOI como condición temática: 86 de los 419 candidatos no tienen uno. Para reproducir los resultados usamos los archivos conservados. Una nueva descarga requiere una llave propia y puede cambiar los conteos.')

heading('5 Cómo se fue afinando la selección')
para('El primer cribado distinguió incluidos, excluidos y casos dudosos. Luego revisamos esos dudosos con reglas y lecturas dirigidas asistidas por IA. Después de los ejercicios humanos ajustamos el criterio y volvimos a separar los casos cuya evidencia era insuficiente. En la última revisión registramos una propuesta para todos ellos, manteniendo el historial.')
figure('02_versiones_del_cribado.png', 'Figura 2. Cambios de propuesta sobre los mismos 4.735 registros. En el primer cribado se agrupan 181 incluidos fuertes y 212 secundarios. La variación de los conteos refleja cambios de criterio y revisión; no mide por sí sola la exactitud.', width=6.65)
table(['Clasificación anterior', 'IN actuales', 'OUT actuales', 'Total'], [
    ['IN', '362', '2', '364'],
    ['OUT', '0', '1.972', '1.972'],
    ['TC', '57', '2.342', '2.399'],
    ['Total actual', '419', '4.316', '4.735']], widths=[2.3, 1.5, 1.5, 1.6])
para('Dos fichas antes incluidas, H003 y H015, pasaron a OUT de acuerdo con las respuestas humanas registradas. Los otros 362 IN se conservaron; no tuvieron una nueva lectura semántica individual en esta última etapa.')
para('Los 419 candidatos reúnen esos 362 registros, 56 casos nuevos con evidencia en título o resumen y un caso incorporado por el juicio humano previo. Las 57 nuevas inclusiones tuvieron lectura directa con IA. El procesamiento del conjunto completo se distingue de esas lecturas dirigidas.')

heading('6 Criterio de desempeño y resolución de dudas')
para('Para proponer la inclusión en el subconjunto de desempeño buscamos una relación concreta entre la tarea, su evaluación y un resultado operacional. El criterio quedó organizado en tres requisitos.')
table(['Requisito', 'Qué buscamos'], [
    ['A Contexto y tarea', 'Equipo o proceso reconocible en minería a cielo abierto.'],
    ['B Evaluación propia', 'Caracterización, comparación, intervención, modelo o simulación del equipo o tarea.'],
    ['C Resultado operacional', 'Producción, tiempo, esperas, disponibilidad, utilización, consumo, costo, distancia, frecuencia o caudal ligados a esa evaluación.']], widths=[1.7, 5.2])
para('Las actividades auxiliares pueden incluirse por su propio resultado, como caudal de bombeo o eficiencia del riego. También admitimos resultados modelados y caracterizaciones de partida. Una propuesta de mejora, la precisión de un predictor o la resistencia de una pieza aislada no bastan para acreditar el desempeño de la tarea.')
figure('03_cierre_de_TC.png', 'Figura 3. Decisiones propuestas para los 2.399 casos antes pendientes. Las exclusiones documentales indican que la ficha no permitió justificar su inclusión.', width=6.65)
para('Se propusieron 57 inclusiones y 2.342 exclusiones. De estas últimas, 616 corresponden a un foco distinto del operacional y 1.726 a limitaciones de la evidencia: 1.694 no acreditan suficientemente un resultado de desempeño, 22 presentan una fuente no fiable y 10 no permiten confirmar el contexto.')
para('El código procesó todas las fichas y la IA realizó 118 lecturas directas de los antiguos dudosos. No se leyeron los 2.399 textos completos. Un OUT por falta de evidencia deja el artículo fuera del subconjunto actual, pero no demuestra que el texto completo sea irrelevante. El criterio se fue ajustando durante el trabajo y ese cambio debe quedar explicado en el método.')

heading('7 Publicaciones por año')
para('Contamos los años de publicación en la base amplia y en los 419 candidatos de desempeño. Cada panel usa su propia escala y población. La serie principal llega hasta 2025; dejamos 2026 separado porque el año todavía está incompleto.')
figure('04_publicaciones_por_ano.png', 'Figura 4. Distribución de publicaciones recuperadas entre 2000 y 2025. Fuente: año de publicación registrado en la matriz. Los paneles corresponden a la base amplia y al subconjunto de desempeño.', width=6.65)
table(['Cobertura', 'Base amplia', 'Subconjunto de desempeño'], [
    ['2000 a 2025', '3.784', '339'],
    ['Antes de 2000', '560', '22'],
    ['2026 incompleto', '382', '58'],
    ['Sin año', '9', '0'],
    ['Total', '4.735', '419']], widths=[2.7, 1.8, 2.4])
para('En 2025 aparecen 432 publicaciones de la base amplia y 44 del subconjunto. La serie describe lo recuperado por nuestra consulta. La cobertura de OpenAlex y los filtros influyen en ella, por lo que el aumento observado no se interpreta directamente como crecimiento equivalente de toda la investigación minera.')

heading('8 Procesos y actividades')
para('Las etiquetas de proceso se obtuvieron automáticamente y se conservaron para comparar ambas bases. Una publicación puede mencionar varias actividades; se cuenta una vez dentro de cada etiqueta. Por eso no corresponde sumar las barras como si fueran grupos separados.')
figure('05_menciones_de_procesos.png', 'Figura 5. Menciones de procesos en la base amplia de 4.735 registros y el subconjunto de 419. Se usan escalas distintas. La presencia de una palabra no confirma que ese proceso sea el objeto central del artículo.', width=6.75)
para('En la base amplia aparecen 2.051 referencias con vocabulario de geotecnia o caracterización, 802 de tronadura, 437 de perforación y 413 de transporte o acarreo. Según el esquema del profesor, buena parte de geotecnia y caracterización deberá revisarse dentro de captura y procesamiento de información.')
para('Entre los 419 candidatos se detectan 110 menciones de transporte o acarreo, 75 de tronadura, 50 de perforación y 35 de carguío. Estos conteos todavía deben completarse con la clasificación del problema que estudia cada artículo: planificación, asignación o ejecución.')
para('Desagüe y drenaje aparece en 308 referencias amplias. Aun así, algunos estudios elegibles de bombeo no reciben esa etiqueta. Antes de concluir que hay poca investigación en una actividad auxiliar necesitamos revisar sinónimos, idioma y menciones incidentales.')

heading('9 Equipos presentes en la literatura')
para('El análisis por equipos complementa el análisis por procesos. Los diccionarios agrupan algunas familias y pueden detectar menciones de contexto. Un artículo que nombra un camión o una pala no necesariamente evalúa ese equipo.')
figure('06_menciones_de_equipos.png', 'Figura 6. Menciones automáticas de equipos en ambas poblaciones. Las familias pueden superponerse; se mantiene la agrupación original de chancadores o correas.', width=6.75)
table(['Coincidencia solo en el título', 'Base amplia', 'Subconjunto'], [
    ['Camión de extracción', '260', '81'],
    ['Pala o excavadora', '155', '48'],
    ['Chancador o correa', '84', '21']], widths=[3.5, 1.5, 1.9])
para('Al restringir la búsqueda al título cambian los recuentos. Esto ayuda a distinguir una posible investigación centrada en el equipo de una mención más general, aunque el título tampoco basta para decidir.')
para('Los valores bajos de aljibes, motoniveladoras y bombas son puntos de revisión. Pueden reflejar diferencias de vocabulario o artículos descritos por la función que cumplen. Con estos conteos todavía no podemos dar por confirmada la hipótesis de escasa literatura sobre auxiliares.')

heading('10 Enfoques de estudio y revisión humana')
figure('07_focos_en_cada_base.png', 'Figura 7. Porcentaje de referencias con cada etiqueta temática dentro de su propia base. Son menciones automáticas; las diferencias también dependen del criterio usado para seleccionar el subconjunto.', width=6.65)
para('La etiqueta de desempeño aparece en 1.256 referencias amplias y 281 del subconjunto. Detectar ese vocabulario es distinto de comprobar una evaluación operacional. Además, el análisis de producción científica y citas propio de la bibliometría debe distinguirse del desempeño de los equipos y tareas que buscamos estudiar.')
figure('08_calibracion_humana.png', 'Figura 8. Decisiones humanas registradas en dos ejercicios con sugerencias visibles. Se conservan las reconsideraciones IN de P04 y P07 y la respuesta histórica TC de H011.', width=6.65)
para('El primer piloto quedó en 8 IN y 12 OUT; el segundo lote, en 6 IN, 13 OUT y 1 TC. Son 40 fichas distintas. En el segundo lote hubo 15 acuerdos y 5 diferencias con la sugerencia previa. Estos ejercicios sirvieron para ajustar el criterio y no son una validación independiente.')
para('H013 se mantiene incluido por la respuesta humana, con la discrepancia de IA registrada aparte. H003 y H015 quedan fuera según sus respuestas humanas. H011 conserva su TC humano y recibe una nueva sugerencia IA de inclusión por el volumen de bombeo evaluado; no se afirma que el software propuesto haya sido validado.')

heading('11 Código y reproducción de los resultados')
para('El paquete reúne los códigos de descarga y filtrado, las entradas conservadas, las auditorías y los resultados. La reproducción local parte de los 9.534 registros reconstruidos y permite repetir el filtro amplio y las revisiones posteriores. No necesita hacer otra consulta a OpenAlex.')
table(['Etapa', 'Script', 'Salida principal'], [
    ['Extracción', '01_extraer_openalex.py', 'Universo bibliográfico'],
    ['Filtro general', '02_filtrar_y_unificar.py', 'Base amplia e historial'],
    ['Primer cribado', '06_screening_multidimensional.py', 'Matriz operacional'],
    ['Revisión de dudas', '11_resolver_dudosos.py', 'Decisiones por ficha'],
    ['Criterio ajustado', '15_aplicar_criterio_calibrado.py', 'Matriz con casos pendientes'],
    ['Última revisión', '22_cerrar_TC.py', 'Matriz y particiones actuales']], widths=[1.15, 3.65, 2.1])
para('Para repetir el procesamiento basta Python 3.10 o superior. Desde la carpeta raíz del paquete, en PowerShell:')
code('py -3 00_reproducir_local.py')
para('La ejecución escribe en una carpeta nueva y conserva los archivos originales. Se comprobó la cadena completa: sus 16 salidas estables coincidieron con las archivadas mediante SHA-256. Las fechas de una ejecución nueva pueden cambiar, pero no alteran ese cotejo.')
para('Las auditorías IA y las respuestas humanas se usan como entradas registradas. El comando reproduce la aplicación de reglas y la integración de esas decisiones; no genera nuevas lecturas con IA ni descarga nuevamente los datos de septiembre.')
para('El README y la guía técnica explican las carpetas, los campos y la consulta completa. Los códigos para regenerar los gráficos y este avance están separados de los motores de selección. Una extracción nueva requiere una llave propia y debe tratarse como otra instantánea.')

heading('12 Trabajo que sigue para el artículo')
para('El cierre de los casos pendientes permite avanzar con una propuesta completa. Ahora falta comprobar su calidad y terminar la clasificación necesaria para responder las preguntas del profesor.')
doc.add_heading('Completar la clasificación temática', level=2)
para('En la base amplia separaremos procesos principales, auxiliares, mantenimiento e información, y registraremos el problema abordado: diseño, planificación, asignación o ejecución. También habrá que distinguir el equipo estudiado de una mención incidental. Con estos cruces podremos contrastar las tres hipótesis y después describir en detalle la literatura sobre auxiliares.')
doc.add_heading('Revisar una muestra independiente', level=2)
para('Seguiremos la indicación del profesor de reducir el corpus y revisar después una muestra humana. La muestra debe incluir candidatos IN, excluidos por foco y excluidos por falta de evidencia. También necesitamos un control separado de los 4.799 que quedaron fuera del filtro amplio, para revisar posibles omisiones.')
para('Hay 260 identificadores usados en desarrollo, calibración o lecturas dirigidas documentadas. Quedan 4.475 referencias amplias fuera de ese conjunto. De la muestra anterior de 180, 159 no figuran en esa exposición registrada. El diseño y los pesos deben actualizarse al criterio actual antes de usarla para estimar errores.')
doc.add_heading('Preparar los análisis bibliométricos', level=2)
para('Después de depurar los términos podremos construir redes de coocurrencia y comparar perfiles temáticos. Para co-citación o acoplamiento bibliográfico faltan las listas de referencias enlazadas: el número total de citas no permite reconstruir esos vínculos. Una red de colaboración también requerirá recuperar identificadores de autores e instituciones.')
para('La revisión desde computación puede ayudar a comprobar las reglas, sinónimos, deduplicación y manejo de versiones. Desde el criterio minero debemos precisar las tareas, equipos e indicadores. Los cambios quedarán registrados por ID, conservando las decisiones anteriores.')
para('Hay 58 alertas bibliográficas heredadas, una dentro de los 419 candidatos. Su revisión debe separarse de la pertinencia temática. Las conclusiones sobre brechas de auxiliares se escribirán después de comprobar el foco de los artículos y la cobertura de la búsqueda.')

heading('13 Fuentes y registro del método')
para('El marco de investigación considera la operación a cielo abierto como un sistema de procesos y propone profundizar en los auxiliares. Las preguntas ordenan ese análisis y las hipótesis se mantienen como afirmaciones por contrastar.')
doc.add_heading('Referencias metodológicas', level=2)
para('Donthu, N., Kumar, S., Mukherjee, D., Pandey, N. y Lim, W. M. (2021). How to conduct a bibliometric analysis: An overview and guidelines. Journal of Business Research, 133, 285-296. DOI: 10.1016/j.jbusres.2021.04.070.')
para('Aria, M. y Cuccurullo, C. (2017). bibliometrix: An R-tool for comprehensive science mapping analysis. Journal of Informetrics, 11(4), 959-975. DOI: 10.1016/j.joi.2017.08.007.')
para('Estas referencias orientan la definición del alcance, recopilación, depuración y análisis bibliométrico. El artículo de royalties compartido sirve como ejemplo de combinación entre bibliometría y análisis de contenido. Las reglas y el apoyo de IA se describirán como parte del método desarrollado para este proyecto.')
doc.add_heading('Documentación de OpenAlex', level=2)
para('Documentación oficial consultada el 8 de octubre de 2026. En este avance se trabajó con los archivos conservados, sin una descarga nueva.')
para('https://help.openalex.org/api/searching/\nhttps://help.openalex.org/api/authentication/\nhttps://help.openalex.org/api/paging/')
doc.add_heading('Qué queda registrado para el artículo', level=2)
table(['Registro', 'Contenido conservado'], [
    ['Procedencia de datos', 'Consulta, fechas, tamaños, archivos y límites de recuperación.'],
    ['Historia por publicación', 'Identificador, propuestas anteriores, decisión actual, motivo y evidencia.'],
    ['Revisión humana e IA', 'Respuestas humanas y sugerencias en campos separados; diferencias documentadas.'],
    ['Reproducibilidad', 'Códigos, entradas, auditorías y comprobaciones de igualdad.']], widths=[2.0, 4.9])
para('Los nombres técnicos de archivos y campos se mantienen en el código y en la guía para conservar la trazabilidad. El registro exacto del modelo y los prompts de todas las etapas anteriores no está completo; esa información faltante no se reconstruye retrospectivamente.')
para('Los mapas y el notebook adjuntos sobre minerales críticos corresponden a otro estudio. Se usan como ejemplos metodológicos y no como resultados de esta base.')

target = OUT / 'Sintesis_IPRE_actualizada_reunion_20261008.docx'
doc.save(target)
print(target)
